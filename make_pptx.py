#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate PPTX: Volcano-scheduled sLLM workload + Prometheus/Grafana monitoring."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- theme ----
NAVY   = RGBColor(0x0B, 0x1F, 0x3A)
BLUE   = RGBColor(0x1E, 0x5A, 0xA8)
ACCENT = RGBColor(0xF2, 0x6D, 0x21)   # volcano orange
LIGHT  = RGBColor(0xF4, 0xF6, 0xFA)
GRAY   = RGBColor(0x55, 0x5B, 0x66)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x1B, 0x9E, 0x57)
FONT   = "Malgun Gothic"   # Windows Korean font

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set(run, size, bold=False, color=NAVY, font=FONT):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = color; run.font.name = font


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    return sp


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt,size,bold,color)."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for (txt, size, bold, color) in para:
            r = p.add_run(); r.text = txt; _set(r, size, bold, color)
    return tb


def header(slide, title, kicker=None, page=None):
    box(slide, 0, 0, SW, Inches(1.15), fill=NAVY)
    box(slide, 0, Inches(1.15), SW, Inches(0.06), fill=ACCENT)
    text(slide, Inches(0.55), Inches(0.18), Inches(11.5), Inches(0.9),
         ([[(kicker, 12, True, ACCENT)]] if kicker else []) +
         [[(title, 26, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    if page:
        text(slide, Inches(12.2), Inches(0.35), Inches(0.9), Inches(0.5),
             [[(page, 12, True, WHITE)]], align=PP_ALIGN.RIGHT)


def bullet(slide, x, y, w, h, items, size=15, gap=8):
    """items: list of (text, level, color, bold)."""
    runs = []
    for it in items:
        t, lvl = it[0], it[1]
        color = it[2] if len(it) > 2 else NAVY
        bold = it[3] if len(it) > 3 else False
        mark = "■ " if lvl == 0 else ("– " if lvl == 1 else "· ")
        indent = "    " * lvl
        runs.append([(indent + mark, size, True, ACCENT if lvl == 0 else GRAY),
                     (t, size, bold, color)])
    return text(slide, x, y, w, h, runs, space_after=gap, line_spacing=1.05)


# ============================================================ Slide 1: Title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=NAVY)
box(s, 0, Inches(5.0), SW, Inches(0.08), fill=ACCENT)
text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(2.2),
     [[("Volcano 기반 sLLM 워크로드", 40, True, WHITE)],
      [("스케줄링 & 실시간 모니터링 구축", 40, True, WHITE)]], space_after=6)
text(s, Inches(0.95), Inches(3.7), Inches(11), Inches(0.8),
     [[("SmolLM 135M 추론 작업을 Volcano로 스케줄링하고 Prometheus·Grafana로 관측", 17, False, RGBColor(0xC8,0xD4,0xE6))]])
text(s, Inches(0.95), Inches(5.25), Inches(11), Inches(1.2),
     [[("KETI AI Storage System  ·  Kubernetes 1.32  ·  Volcano Scheduler", 14, True, RGBColor(0x9F,0xB3,0xD0))],
      [("작성일: 2026-06-02", 13, False, RGBColor(0x9F,0xB3,0xD0))]], space_after=4)

# ============================================================ Slide 2: 한눈에 보기
s = prs.slides.add_slide(BLANK)
header(s, "한눈에 보기 — 무엇을, 어떻게, 어떻게 봤나", kicker="OVERVIEW", page="2")
cards = [
    ("① 무엇을 (What)", ["SmolLM 135M", "초소형 LLM (135M 파라미터)", "Ollama 런타임, CPU 추론"], BLUE),
    ("② 어떻게 (How)", ["Volcano Job 으로 제출", "default 큐 → 마스터 노드 배치", "프롬프트 추론 1회 실행"], ACCENT),
    ("③ 어떻게 봤나 (Monitor)", ["Prometheus 가 메트릭 수집", "Grafana 'Volcano Scheduler'", "대시보드로 실시간 관측"], GREEN),
]
cw = Inches(3.95); gap = Inches(0.3); x0 = Inches(0.55); y0 = Inches(1.55)
for i, (title, lines, c) in enumerate(cards):
    x = x0 + i * (cw + gap)
    box(s, x, y0, cw, Inches(4.6), fill=LIGHT)
    box(s, x, y0, cw, Inches(0.75), fill=c)
    text(s, x, y0, cw, Inches(0.75), [[(title, 16, True, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.25), y0 + Inches(1.0), cw - Inches(0.5), Inches(3.4),
         [[(lines[0], 19, True, NAVY)]] + [[("• " + l, 14, False, GRAY)] for l in lines[1:]],
         space_after=12, line_spacing=1.1)
text(s, Inches(0.55), Inches(6.5), Inches(12.2), Inches(0.7),
     [[("결과: ", 15, True, ACCENT), ("Volcano 스케줄 → 파드 Running → 추론 완료 → 대시보드에 스케줄링 지표 유입 (p99 지연 9.84 ms) 까지 전 구간 검증 완료", 15, False, NAVY)]])

# ============================================================ Slide 3: 아키텍처
s = prs.slides.add_slide(BLANK)
header(s, "시스템 구성 / 데이터 흐름", kicker="ARCHITECTURE", page="3")

def node(x, y, w, h, title, sub, fill, tcol=WHITE):
    box(s, x, y, w, h, fill=fill)
    text(s, x, y, w, h,
         [[(title, 15, True, tcol)]] + ([[(sub, 11, False, tcol)]] if sub else []),
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)

def arrow(x, y, w):
    from pptx.enum.shapes import MSO_SHAPE
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.32))
    a.fill.solid(); a.fill.fore_color.rgb = ACCENT; a.line.fill.background(); a.shadow.inherit = False

y1 = Inches(1.7)
node(Inches(0.55), y1, Inches(2.7), Inches(1.3), "Volcano Job\n(sllm-smollm-demo)", "SmolLM 135M / Ollama", BLUE)
arrow(Inches(3.35), y1 + Inches(0.49), Inches(0.7))
node(Inches(4.15), y1, Inches(2.7), Inches(1.3), "Volcano Scheduler", "default 큐 → master 배치", NAVY)
arrow(Inches(6.95), y1 + Inches(0.49), Inches(0.7))
node(Inches(7.75), y1, Inches(2.6), Inches(1.3), "Pod (Running)", "ai-storage-master / CPU", RGBColor(0x37,0x4A,0x63))
node(Inches(10.55), y1, Inches(2.2), Inches(1.3), "추론 출력", "K8s 설명 텍스트 생성", GREEN)

# monitoring row
y2 = Inches(3.9)
node(Inches(0.55), y2, Inches(3.4), Inches(1.3), "volcano-scheduler :8080\nvolcano-controllers :8081", "/metrics 노출 (27종)", RGBColor(0x37,0x4A,0x63))
arrow(Inches(4.05), y2 + Inches(0.49), Inches(0.7))
node(Inches(4.85), y2, Inches(3.2), Inches(1.3), "Prometheus", "service-endpoints 스크랩", BLUE)
arrow(Inches(8.15), y2 + Inches(0.49), Inches(0.7))
node(Inches(8.95), y2, Inches(3.8), Inches(1.3), "Grafana", "Volcano Scheduler 대시보드\nNodePort 30030", GREEN)

# down arrow connecting pod -> metrics
from pptx.enum.shapes import MSO_SHAPE
da = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(2.0), y1 + Inches(1.35), Inches(0.3), Inches(0.5))
da.fill.solid(); da.fill.fore_color.rgb = GRAY; da.line.fill.background(); da.shadow.inherit = False
text(s, Inches(0.55), Inches(5.55), Inches(12.2), Inches(1.4),
     [[("기반 인프라: ", 13, True, ACCENT), ("Kubernetes 1.32 · CNI = Flannel(10.244.0.0/16) · 단일 마스터(GPU 없음, CPU 추론) · 그라파나 설정은 PVC(nfs-client 5Gi)로 영속화", 13, False, GRAY)]])

# ============================================================ Slide 4: sLLM 모델
s = prs.slides.add_slide(BLANK)
header(s, "sLLM 모델 — SmolLM 135M", kicker="WHAT · 무엇을", page="4")
box(s, Inches(0.55), Inches(1.5), Inches(5.9), Inches(5.3), fill=LIGHT)
text(s, Inches(0.8), Inches(1.7), Inches(5.4), Inches(0.6), [[("모델 사양", 18, True, BLUE)]])
bullet(s, Inches(0.8), Inches(2.4), Inches(5.4), Inches(4.2), [
    ("이름: SmolLM 135M (HuggingFace)", 0, NAVY, True),
    ("파라미터: 1.35억 (135 Million)", 1),
    ("런타임: Ollama (llama.cpp 백엔드)", 1),
    ("모델 크기: 약 91 MB (양자화)", 1),
    ("추론 장치: CPU (Intel Xeon E5-2620 v4)", 1),
    ("GPU 불필요 → 단일 노드에서 즉시 실행", 1, GREEN, True),
    ("Fallback: qwen2.5:0.5b 자동 전환 설정", 1),
])
box(s, Inches(6.7), Inches(1.5), Inches(6.05), Inches(5.3), fill=NAVY)
text(s, Inches(6.95), Inches(1.7), Inches(5.6), Inches(0.6), [[("왜 이 모델인가", 18, True, ACCENT)]])
bullet2 = [
    ("워커 노드 다운 → 마스터 CPU만 가용", 0, WHITE, False),
    ("GPU 없이도 수초 내 추론 가능한 초소형 모델 필요", 1, RGBColor(0xC8,0xD4,0xE6)),
    ("Ollama 로 모델 pull·서빙·추론을 한 컨테이너에서 처리", 0, WHITE, False),
    ("이미지 1개로 재현성·이식성 확보", 1, RGBColor(0xC8,0xD4,0xE6)),
    ("목적: 실제 AI 워크로드로 Volcano 스케줄링을 발생시켜", 0, WHITE, False),
    ("모니터링 파이프라인을 end-to-end 로 검증", 1, RGBColor(0xC8,0xD4,0xE6)),
]
# render white bullets manually
runs = []
for it in bullet2:
    t, lvl, col = it[0], it[1], it[2]
    b = it[3] if len(it) > 3 else False
    mark = "■ " if lvl == 0 else "– "
    runs.append([("    " * lvl + mark, 15, True, ACCENT if lvl==0 else RGBColor(0x9F,0xB3,0xD0)),
                 (t, 15, b, col)])
text(s, Inches(6.95), Inches(2.4), Inches(5.6), Inches(4.2), runs, space_after=10, line_spacing=1.1)

# ============================================================ Slide 5: 어떻게 (Volcano Job)
s = prs.slides.add_slide(BLANK)
header(s, "어떻게 — Volcano Job 제출 & 스케줄링", kicker="HOW · 어떻게", page="5")
text(s, Inches(0.55), Inches(1.45), Inches(6.0), Inches(0.5), [[("실행 단계", 18, True, BLUE)]])
steps = [
    "Volcano Job(vcjob) 매니페스트 작성 — queue: default, minAvailable: 1",
    "schedulerName: volcano 로 Volcano 스케줄러에 위임",
    "Volcano 가 노드 선택 → 파드에 PodGroup 부여",
    "kubelet 이 Ollama 이미지 pull → 모델(91MB) 다운로드",
    "ollama run 으로 프롬프트 추론 1회 실행 → 출력 생성",
    "Job 상태: Running → Completed",
]
runs = [[("%d. " % (i+1), 15, True, ACCENT), (t, 15, False, NAVY)] for i, t in enumerate(steps)]
text(s, Inches(0.55), Inches(2.05), Inches(6.1), Inches(4.6), runs, space_after=13, line_spacing=1.05)

# code box
box(s, Inches(6.85), Inches(1.45), Inches(5.9), Inches(5.4), fill=RGBColor(0x1C,0x24,0x33))
code = [
    "apiVersion: batch.volcano.sh/v1alpha1",
    "kind: Job",
    "metadata: { name: sllm-smollm-demo }",
    "spec:",
    "  schedulerName: volcano",
    "  queue: default",
    "  minAvailable: 1",
    "  tasks:",
    "  - replicas: 1",
    "    name: inference",
    "    template:",
    "      spec:",
    "        nodeSelector:           # 마스터 고정",
    "          kubernetes.io/hostname: ai-storage-master",
    "        containers:",
    "        - name: ollama",
    "          image: ollama/ollama:latest",
    "          args: [ ollama pull smollm:135m;",
    "                  ollama run smollm:135m \"...\" ]",
    "          resources:",
    "            requests: { cpu: 2, memory: 2Gi }",
]
tb = s.shapes.add_textbox(Inches(7.05), Inches(1.6), Inches(5.6), Inches(5.1)); tf = tb.text_frame
tf.word_wrap = True
for i, ln in enumerate(code):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(1)
    r = p.add_run(); r.text = ln
    r.font.size = Pt(10.5); r.font.name = "Consolas"
    r.font.color.rgb = RGBColor(0x7E,0xE0,0x8A) if ("#" in ln or ln.strip().endswith(":")) else RGBColor(0xE6,0xEC,0xF2)

# ============================================================ Slide 6: 트러블슈팅
s = prs.slides.add_slide(BLANK)
header(s, "스케줄링 트러블슈팅 포인트", kicker="HOW · 문제해결", page="6")
box(s, Inches(0.55), Inches(1.55), Inches(5.95), Inches(4.9), fill=RGBColor(0xFD,0xEE,0xE6))
text(s, Inches(0.8), Inches(1.75), Inches(5.5), Inches(0.5), [[("⚠ 문제", 18, True, ACCENT)]])
bullet(s, Inches(0.8), Inches(2.45), Inches(5.5), Inches(3.8), [
    ("첫 제출 시 파드가 죽은 워커에 배치됨", 0, NAVY, True),
    ("ai-storage-worker-01 = NotReady(unreachable)", 1),
    ("tolerations: {operator: Exists} 가", 0, NAVY, True),
    ("unreachable taint 까지 모두 허용 → 죽은 노드 선택", 1),
    ("결과: 파드가 6분 넘게 Pending", 1, ACCENT, True),
])
box(s, Inches(6.8), Inches(1.55), Inches(5.95), Inches(4.9), fill=RGBColor(0xE7,0xF5,0xEC))
text(s, Inches(7.05), Inches(1.75), Inches(5.5), Inches(0.5), [[("✔ 해결", 18, True, GREEN)]])
bullet(s, Inches(7.05), Inches(2.45), Inches(5.5), Inches(3.8), [
    ("광범위 toleration 제거", 0, NAVY, True),
    ("nodeSelector 로 마스터에 고정", 0, NAVY, True),
    ("kubernetes.io/hostname: ai-storage-master", 1),
    ("Job 재제출 → 마스터에 정상 배치", 0, GREEN, True),
    ("파드 Running, 추론 정상 수행", 1),
])
text(s, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.6),
     [[("교훈: ", 14, True, ACCENT), ("노드 일부가 다운된 환경에서는 toleration 범위를 넓히지 말고 nodeSelector/affinity 로 가용 노드를 명시할 것", 14, False, GRAY)]])

# ============================================================ Slide 7: 모니터링 구성
s = prs.slides.add_slide(BLANK)
header(s, "어떻게 모니터링했나 — Prometheus · Grafana", kicker="MONITOR · 관측", page="7")
text(s, Inches(0.55), Inches(1.45), Inches(6), Inches(0.5), [[("① Prometheus ← Volcano", 17, True, BLUE)]])
bullet(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(2.4), [
    ("scheduler(:8080)·controller(:8081) 스크랩", 0),
    ("job=kubernetes-service-endpoints, 둘 다 up", 1),
    ("volcano_* 메트릭 27종 수집", 0, GREEN, True),
    ("podgroup 상태·큐 할당량·스케줄 지연 등", 1),
])
text(s, Inches(0.55), Inches(4.35), Inches(6), Inches(0.5), [[("② Grafana ← Prometheus", 17, True, GREEN)]])
bullet(s, Inches(0.55), Inches(4.9), Inches(6.0), Inches(2.2), [
    ("Helm 으로 datasource·대시보드 provisioning", 0),
    ("'Volcano Scheduler' 대시보드 8개 패널", 0, NAVY, True),
    ("PVC(nfs-client 5Gi)로 설정 영속화", 1, GREEN, True),
    ("접속: http://10.0.4.80:30030 (NodePort)", 1),
])
# right: dashboard panels list
box(s, Inches(6.95), Inches(1.5), Inches(5.8), Inches(5.4), fill=LIGHT)
text(s, Inches(7.2), Inches(1.7), Inches(5.3), Inches(0.5), [[("대시보드 패널 구성", 16, True, BLUE)]])
panels = [
    "Running / Pending / Inqueue PodGroups",
    "Preemption Victims (총합)",
    "PodGroups by State (큐별 시계열)",
    "Queue Allocated CPU (milli)",
    "Queue Allocated Memory (bytes)",
    "E2E Scheduling Latency p50/p90/p99",
]
runs = [[("▪ ", 13, True, ACCENT), (p, 13.5, False, NAVY)] for p in panels]
text(s, Inches(7.2), Inches(2.4), Inches(5.3), Inches(4.2), runs, space_after=14, line_spacing=1.1)

# ============================================================ Slide 8: 결과
s = prs.slides.add_slide(BLANK)
header(s, "결과 — 추론 출력 & 관측 지표", kicker="RESULT", page="8")
# metrics cards
mc = [("Running PodGroup", "1 → 0", "실행 중 1, 완료 후 0"),
      ("E2E 스케줄 지연 p99", "9.84 ms", "히스토그램에 실데이터 유입"),
      ("Volcano 메트릭", "27 종", "Prometheus 수집 확인"),
      ("Job 상태", "Completed", "추론 정상 종료")]
cw = Inches(2.95); x0 = Inches(0.55); y0 = Inches(1.55)
for i, (t, v, sub) in enumerate(mc):
    x = x0 + i * (cw + Inches(0.13))
    box(s, x, y0, cw, Inches(1.7), fill=LIGHT)
    box(s, x, y0, Inches(0.12), Inches(1.7), fill=ACCENT)
    text(s, x + Inches(0.2), y0 + Inches(0.12), cw - Inches(0.3), Inches(1.5),
         [[(t, 12.5, True, GRAY)], [(v, 26, True, NAVY)], [(sub, 10.5, False, GRAY)]],
         space_after=3)
# inference output
box(s, Inches(0.55), Inches(3.6), Inches(12.2), Inches(3.25), fill=RGBColor(0x1C,0x24,0x33))
text(s, Inches(0.8), Inches(3.75), Inches(11.7), Inches(0.5),
     [[("추론 출력 (프롬프트: \"In two sentences, explain what Kubernetes is\")", 14, True, ACCENT)]])
out = ('"Kubernetes (also known as K8s) is a popular open-source container orchestration '
       'application and the operating system or infrastructure it runs on. It enables '
       'autonomous, modular and scalable deployment of applications in a controlled and '
       'automated manner..."')
text(s, Inches(0.8), Inches(4.4), Inches(11.7), Inches(2.3),
     [[(out, 14, False, RGBColor(0xE6,0xEC,0xF2))],
      [("※ 135M 초소형 모델 특성상 다소 장황하나, CPU 추론으로 정상적인 텍스트 생성 확인", 11.5, False, RGBColor(0x9F,0xB3,0xD0))]],
     space_after=10, line_spacing=1.15)

# ============================================================ Slide 9: 마무리
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, fill=NAVY)
box(s, Inches(0.9), Inches(1.3), Inches(0.14), Inches(1.7), fill=ACCENT)
text(s, Inches(1.25), Inches(1.3), Inches(11), Inches(1.8),
     [[("정리 & 다음 단계", 32, True, WHITE)]])
summ = [
    ("검증 완료된 전 구간 파이프라인", 0, ACCENT, True),
    ("Volcano Job(SmolLM 135M) → 스케줄 → CPU 추론 → Completed", 1, WHITE, False),
    ("Prometheus 메트릭 수집 → Grafana 대시보드 실시간 표시", 1, WHITE, False),
    ("다음 단계 제안", 0, ACCENT, True),
    ("장시간/다중 Job 제출로 큐 경합·지연 그래프 강화", 1, WHITE, False),
    ("Deployment 상시 서빙으로 라이브 부하 모니터링", 1, WHITE, False),
    ("워커 노드 복구 후 멀티노드 스케줄링 검증", 1, WHITE, False),
]
runs = []
for t, lvl, col, b in summ:
    mark = "■ " if lvl == 0 else "– "
    runs.append([("    " * lvl + mark, 17, True, ACCENT if lvl==0 else RGBColor(0x9F,0xB3,0xD0)),
                 (t, 17, b, col)])
text(s, Inches(1.25), Inches(3.2), Inches(11), Inches(3.2), runs, space_after=12, line_spacing=1.15)
text(s, Inches(1.25), Inches(6.6), Inches(11), Inches(0.6),
     [[("대시보드: http://10.0.4.80:30030  →  Volcano 폴더  →  'Volcano Scheduler'", 13, True, RGBColor(0xC8,0xD4,0xE6))]])

out_path = "/root/workspace/volcano-sllm-monitoring.pptx"
prs.save(out_path)
print("saved:", out_path, "| slides:", len(prs.slides._sldIdLst))

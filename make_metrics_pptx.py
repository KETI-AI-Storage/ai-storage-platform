#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate PPTX: Volcano 31개 메트릭 설명 표 (표만)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- theme ----
NAVY   = RGBColor(0x0B, 0x1F, 0x3A)
BLUE   = RGBColor(0x1E, 0x5A, 0xA8)
ACCENT = RGBColor(0xF2, 0x6D, 0x21)   # volcano orange
LIGHT  = RGBColor(0xF4, 0xF6, 0xFA)
ROW2   = RGBColor(0xEA, 0xEF, 0xF7)
GRAY   = RGBColor(0x33, 0x38, 0x40)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = "Malgun Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(24); r.font.bold = True
    r.font.color.rgb = NAVY; r.font.name = FONT


def add_table(slide, headers, rows, col_widths, top=Inches(1.0),
              header_fs=11, body_fs=10, row_h=Inches(0.34)):
    nrows = len(rows) + 1
    ncols = len(headers)
    total_w = sum(col_widths)
    left = (SW - total_w) // 2
    gtbl = slide.shapes.add_table(nrows, ncols, left, top, total_w, row_h * nrows)
    tbl = gtbl.table
    # disable banding style noise; we color manually
    tbl.first_row = False
    tbl.horz_banding = False
    for i, w in enumerate(col_widths):
        tbl.columns[i].width = w
    # header
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.05); cell.margin_right = Inches(0.05)
        cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = h
        r.font.size = Pt(header_fs); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = FONT
    # body
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if ri % 2 else ROW2
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.01); cell.margin_bottom = Inches(0.01)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if ci in (0, 2, 3) else PP_ALIGN.LEFT
            r = p.add_run(); r.text = val
            r.font.size = Pt(body_fs); r.font.name = FONT
            r.font.color.rgb = GRAY
            if ci == 1:
                r.font.bold = True; r.font.color.rgb = NAVY
            if ci == 0:
                r.font.bold = True; r.font.color.rgb = ACCENT
    return gtbl


# column layout: # | 메트릭 | 타입 | 단위 | 의미
HEADERS = ["#", "메트릭 이름", "타입", "단위", "의미 / 무엇을 뜻하는가"]
CW = [Inches(0.45), Inches(4.2), Inches(1.0), Inches(0.95), Inches(6.1)]
ROWH = Inches(0.40)

# ============ Slide 1: 표지 ============
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
bar = s.shapes.add_shape(1, 0, Inches(4.05), SW, Inches(0.12))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
t = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.4))
p = t.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Volcano 메트릭 전수 해설"
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
t2 = s.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(1.0))
p2 = t2.text_frame.paragraphs[0]
r2 = p2.add_run()
r2.text = "scheduler :8080 (25개) + controller :8081 (6개) = 총 31개"
r2.font.size = Pt(18); r2.font.color.rgb = RGBColor(0xC9, 0xD6, 0xEC); r2.font.name = FONT

# ============ Slide 2: A-1 스케줄링 지연/성능 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "A. 스케줄러 — ① 스케줄링 지연 / 성능 (8개)")
rows = [
 ["1", "volcano_e2e_scheduling_latency_milliseconds", "histogram", "ms",
  "스케줄링 1사이클 전체(알고리즘+바인딩) 시간. 가장 핵심 지표"],
 ["2", "volcano_e2e_job_scheduling_latency_milliseconds", "histogram", "ms",
  "Job(=PodGroup, gang 단위) 하나의 end-to-end 스케줄링 시간"],
 ["3", "volcano_task_scheduling_latency_milliseconds", "histogram", "ms",
  "Task(=Pod 1개)를 노드에 배치 결정하는 알고리즘 지연(바인딩 제외)"],
 ["4", "volcano_action_scheduling_latency_milliseconds", "histogram", "ms",
  "액션(단계)별 지연. enqueue/allocate/preempt/backfill 각각의 소요"],
 ["5", "volcano_plugin_scheduling_latency_milliseconds", "histogram", "ms",
  "플러그인별 지연(label: plugin, OnSession). 병목 플러그인 진단"],
 ["6", "volcano_e2e_job_scheduling_duration", "gauge", "ms",
  "특정 Job의 시작~마지막 스케줄링 경과 시간 (last-start)"],
 ["7", "volcano_e2e_job_scheduling_start_time", "gauge", "epoch s",
  "그 Job의 스케줄링 시작 시각(타임스탬프)"],
 ["8", "volcano_e2e_job_scheduling_last_time", "gauge", "epoch s",
  "그 Job이 마지막으로 스케줄링된 시각(재스케줄 시 갱신)"],
]
add_table(s, HEADERS, rows, CW, top=Inches(1.0), row_h=ROWH)

# ============ Slide 3: A-2,A-3 선점/실패 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "A. 스케줄러 — ② 선점(Preemption) + 스케줄 실패 (4개)")
rows = [
 ["9", "volcano_total_preemption_attempts", "counter", "회",
  "클러스터 누적 선점 시도 횟수(낮은 우선순위 Pod 쫓고 자리 확보)"],
 ["10", "volcano_pod_preemption_victims", "gauge", "개",
  "선점으로 실제 쫓겨난(victim) Pod 수"],
 ["11", "volcano_unschedule_task_count", "gauge", "개",
  "스케줄 실패한 Task(Pod) 수 (자원/제약으로 노드 배치 불가)"],
 ["12", "volcano_unschedule_job_count", "gauge", "개",
  "스케줄 실패한 Job 수 (gang 조건 미충족으로 통째 실패)"],
]
add_table(s, HEADERS, rows, CW, top=Inches(1.05), row_h=Inches(0.5))
note = s.shapes.add_textbox(Inches(0.6), Inches(3.6), Inches(12.1), Inches(0.5))
np_ = note.text_frame.paragraphs[0]
nr = np_.add_run(); nr.text = "※ 현재 클러스터: 9·10 모두 0 (선점 미발생)."
nr.font.size = Pt(12); nr.font.italic = True; nr.font.color.rgb = GRAY; nr.font.name = FONT

# ============ Slide 4: A-4 큐 자원회계 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "A. 스케줄러 — ③ 큐 자원 회계: Allocated/Request/Deserved (9개)")
rows = [
 ["13", "volcano_queue_allocated_milli_cpu", "gauge", "mCPU", "큐에 현재 실제 할당된 CPU"],
 ["14", "volcano_queue_allocated_memory_bytes", "gauge", "bytes", "큐에 현재 실제 할당된 메모리"],
 ["15", "volcano_queue_allocated_scalar_resources", "gauge", "개수", "큐 할당 기타자원(GPU/pods/storage, label:resource)"],
 ["16", "volcano_queue_request_milli_cpu", "gauge", "mCPU", "큐 내 Pod들이 요청(request)한 CPU 합계"],
 ["17", "volcano_queue_request_memory_bytes", "gauge", "bytes", "큐 요청 메모리 합계"],
 ["18", "volcano_queue_request_scalar_resources", "gauge", "개수", "큐 요청 기타자원 합계(label:resource)"],
 ["19", "volcano_queue_deserved_milli_cpu", "gauge", "mCPU", "proportion이 계산한 큐가 받을 자격(deserved) CPU 몫"],
 ["20", "volcano_queue_deserved_memory_bytes", "gauge", "bytes", "deserved 메모리 몫"],
 ["21", "volcano_queue_deserved_scalar_resources", "gauge", "개수", "deserved 기타자원 몫(label:resource)"],
]
add_table(s, HEADERS, rows, CW, top=Inches(1.0), row_h=ROWH)
note = s.shapes.add_textbox(Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.5))
np_ = note.text_frame.paragraphs[0]
nr = np_.add_run()
nr.text = "※ 해석축: Allocated(실제받음) ↔ Request(요청) ↔ Deserved(받을자격). allocated>deserved 면 overused."
nr.font.size = Pt(11); nr.font.italic = True; nr.font.color.rgb = GRAY; nr.font.name = FONT

# ============ Slide 5: A-5 공정성 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "A. 스케줄러 — ④ 큐/잡 공정성 지표 (4개)")
rows = [
 ["22", "volcano_queue_weight", "gauge", "-",
  "큐 가중치. deserved 몫 계산 비율(weight 2 = weight 1의 2배 자원)"],
 ["23", "volcano_queue_share", "gauge", "-",
  "큐의 자원 점유 share. share 낮은 큐를 다음 스케줄에서 우선"],
 ["24", "volcano_queue_overused", "gauge", "0/1",
  "큐가 deserved 초과 사용 중이면 1 → 더 이상 새 자원 못 받음"],
 ["25", "volcano_job_share", "gauge", "-",
  "개별 Job의 DRF dominant share. 같은 큐 내 Job 우선순위 결정"],
]
add_table(s, HEADERS, rows, CW, top=Inches(1.05), row_h=Inches(0.55))
note = s.shapes.add_textbox(Inches(0.6), Inches(4.0), Inches(12.1), Inches(0.5))
np_ = note.text_frame.paragraphs[0]
nr = np_.add_run()
nr.text = "※ 현재 default 큐 overused=1 (deserved가 CPU 2core·pods 1로 작게 잡힘) → 큐 capability 점검 포인트."
nr.font.size = Pt(12); nr.font.italic = True; nr.font.color.rgb = ACCENT; nr.font.name = FONT

# ============ Slide 6: B 컨트롤러 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "B. 컨트롤러 (:8081) — Job/PodGroup 라이프사이클 (6개)")
rows = [
 ["26", "volcano_job_completed_phase_count", "counter", "개",
  "Completed 단계 도달한 vcjob 누적 수(label: job_name, queue_name)"],
 ["27", "volcano_queue_pod_group_pending_count", "gauge", "개",
  "큐 내 Pending(큐 승인 전 대기) PodGroup 수"],
 ["28", "volcano_queue_pod_group_inqueue_count", "gauge", "개",
  "Inqueue(큐 승인됨, 스케줄 가능) PodGroup 수"],
 ["29", "volcano_queue_pod_group_running_count", "gauge", "개",
  "Running(멤버 Pod 충분히 떠서 실행 중) PodGroup 수"],
 ["30", "volcano_queue_pod_group_completed_count", "gauge", "개",
  "Completed(작업 종료) PodGroup 수"],
 ["31", "volcano_queue_pod_group_unknown_count", "gauge", "개",
  "Unknown(일부 Pod 비정상 등 상태 불명) PodGroup 수"],
]
add_table(s, HEADERS, rows, CW, top=Inches(1.05), row_h=Inches(0.5))
note = s.shapes.add_textbox(Inches(0.6), Inches(4.4), Inches(12.1), Inches(0.6))
np_ = note.text_frame.paragraphs[0]
nr = np_.add_run()
nr.text = "※ PodGroup 상태머신: Pending → Inqueue → Running → Completed (또는 Unknown). 큐별 실시간 분포."
nr.font.size = Pt(12); nr.font.italic = True; nr.font.color.rgb = GRAY; nr.font.name = FONT

out = "/root/workspace/volcano-metrics-table.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides._sldIdLst))
